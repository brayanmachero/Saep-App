"""Read commercial monthly/hourly tables and additional salary matrix formats."""
import datetime
import io
import json
import re
import sys
import zipfile
from pathlib import Path
import openpyxl
from extraer_matrices import extract_sheet, reference_date, source_context
from preparar_sabana import normal, number

PRICES = {'COSTO FINAL MES': 'mes', 'OFERTA VALOR HORA': 'hora',
          'OFERTA VALOR HORA EXTRA (50%)': 'hora extra 50%',
          'OFERTA VALOR HORA EXTRA FESTIVO (100%)': 'hora extra 100%',
          'PRECIO UNITARIO': None}
TABLE_PRICES = PRICES | {'TARIFA': None, 'HORA NORMAL': 'hora', 'TARIFA DIA': 'día', 'TARIFA HORA': 'hora', 'HHEE 50%': 'hora extra 50%', 'HHEE 100%': 'hora extra 100%'}


def money(value):
    if number(value) is not None:
        return value
    if isinstance(value, str):
        # Only unambiguous grouped Chilean currency text, not decimal guesses.
        value = value.strip().replace('$', '').strip()
        if re.fullmatch(r'\d{1,3}(?:\.\d{3})+(?:,\d+)?', value):
            return float(value.replace('.', '').replace(',', '.'))
    return None


def extract_table(rows, raw_rows, title):
    accepted, issues, header = [], [], None
    for ri, row in enumerate(rows, 1):
        labels = [normal(c.value) for c in row]
        if labels.count('CARGO') == 1 and any(x in TABLE_PRICES for x in labels) and any(x.startswith('SUELDO BASE') for x in labels):
            header = row
            continue
        if header is None:
            continue
        headers = {normal(c.value): c.column for c in header if c.value is not None}
        def get(col):
            return row[col - 1].value if col <= len(row) else None
        cargo = get(headers['CARGO'])
        if not isinstance(cargo, str) or not cargo.strip() or len(cargo.strip()) > 180 or re.search(r'TOTAL|CONSIDERACI|INSTRUCCI|^NOTA|NO USAR|CON ERROR', normal(cargo)):
            continue
        salary = next(c.column for c in header if normal(c.value).startswith('SUELDO BASE'))
        if money(get(salary)) is None:
            continue
        date, period = None, 'Período por confirmar'
        if 'ANO' in headers and 'MES' in headers:
            year, month = get(headers['ANO']), get(headers['MES'])
            m = re.fullmatch(r'P?0?(1[0-2]|[1-9])', str(month or ''))
            if number(year) and m:
                try:
                    date = datetime.date(int(year), int(m[1]), 1).isoformat()
                    period = f'{int(year)} {month}'
                except ValueError:
                    pass
        if 'FECHA' in headers and reference_date(get(headers['FECHA'])):
            date = reference_date(get(headers['FECHA']))
            period = str(get(headers['FECHA']))
        components = []
        for c in header:
            if c.value is None:
                continue
            raw = raw_rows[ri - 1][c.column - 1].value
            components.append(dict(concepto=str(c.value).strip(), valor=get(c.column), celda=f'{openpyxl.utils.get_column_letter(c.column)}{ri}',
                                   formula=raw if isinstance(raw, str) and raw.startswith('=') else None))
        if any(isinstance(c['valor'], str) and c['valor'].startswith('#') for c in components):
            issues.append(dict(hoja=title, fila=ri, motivo='Fila con errores guardados; no importada.'))
            continue
        for h in header:
            key = normal(h.value)
            if key not in TABLE_PRICES:
                continue
            price = money(get(h.column))
            if price is None or price < 0:
                continue
            center_col = headers.get('CD') or headers.get('CENTRO DE COSTO')
            mode_col = headers.get('MODALIDAD')
            conditions = [str(get(headers[n])).strip() for n in ('JORNADA', 'TURNO', 'ZONA', 'CODIGO INTERNO') if n in headers and get(headers[n]) is not None]
            accepted.append(dict(cargo=cargo.strip(), precio=price, unidad=TABLE_PRICES[key],
                                 variante=(' · '.join([str(h.value).strip()] + conditions))[:255],
                                 fecha_referencia=date, periodo=period, hoja=title, celda=row[h.column - 1].coordinate,
                                 conceptos=components, centro_fila=str(get(center_col)).strip() if center_col and get(center_col) else None,
                                 modalidad_fila=str(get(mode_col)).strip() if mode_col and get(mode_col) else None))
    return accepted, issues


def extend(root):
    base = json.loads((root / 'lote-sabana-ampliado.json').read_text(encoding='utf-8'))
    inventory = json.loads((root / 'inventory.json').read_text(encoding='utf-8'))
    sources = {s['clave']: s for s in base['fuentes']}
    seen = {(r['fuente_clave'], r['hoja'], r['celda']) for r in base['tarifas']}
    report, new = [], []
    for item in inventory:
        key = item['archive'] + '/' + item['path']
        source = sources[key]
        if item['archive'] == 'IDL.zip' or source['estado'] == 'excluida' or 'MODELO NUEVO' in normal(item['path']):
            continue
        with zipfile.ZipFile(root / 'sources' / item['archive']) as archive:
            content = archive.read(item['path'])
        values = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        customer, center, mode = source_context(item['path'])
        if {'LTS', 'WALMART'} & set(re.split(r'\W+', normal(item['path']))):
            customer = 'Walmart Chile'  # LTS commercial centers map to this client in the audited master.
        if normal(Path(item['path']).name).startswith('EMERGENT COLD'):
            customer = 'Emergent Cold'
        count, issues = 0, []
        for sheet in values:
            if sheet.max_row > 10000 or re.search(r'LIBRO.*REM|MAESTRO.*EMPLEAD', normal(sheet.title)):
                continue
            rows, raw = list(sheet.iter_rows()), list(formulas[sheet.title].iter_rows())
            matrices, matrix_issues = extract_sheet(rows, raw, sheet.title, PRICES)
            tables, table_issues = extract_table(rows, raw, sheet.title)
            issues.extend(matrix_issues + table_issues)
            for r in matrices + tables:
                label = r.pop('etiqueta_precio', None)
                identity = (key, sheet.title, r['celda'])
                if identity in seen or (label and (key, sheet.title, label) in seen):
                    continue
                seen.add(identity)
                row_center = r.pop('centro_fila', None)
                row_mode = r.pop('modalidad_fila', None)
                r.setdefault('variante', r['unidad'] or 'Precio unitario')
                new.append(dict(fuente_clave=key, cliente=customer, centro=row_center or center, modalidad=row_mode or mode,
                                categoria=None, moneda='CLP', **r,
                                observaciones='Tabla/matriz comercial: valores guardados sin recalcular, unidad según etiqueta explícita; fechas solo de fila o bloque. Cliente/centro de origen, categoría y vigencia sujetos a revisión comercial. Las variantes y conceptos se conservan en el detalle.'))
                count += 1
        values.close()
        formulas.close()
        if count:
            source['estado'] = 'parcial'
            source['observaciones'] = 'Lectura parcial de matrices y tablas por cargo: costo final mensual, precio unitario y tarifas horarias explícitas. Restan estructuras ambiguas y validación comercial; no acredita vigencia.'
        report.append(dict(fuente=key, nuevas=count, incidencias=issues))
    base['tarifas'].extend(new)
    (root / 'lote-sabana-tablas.json').write_text(json.dumps(base, ensure_ascii=False, default=str), encoding='utf-8')
    (root / 'cobertura-tablas.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(dict(nuevas=len(new), fuentes_con_nuevas=sum(x['nuevas'] > 0 for x in report), total=len(base['tarifas']))))


if __name__ == '__main__':
    extend(Path(sys.argv[1]))
