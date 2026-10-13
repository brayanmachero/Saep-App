"""Read-only extraction of the audited sources. Never infer commercial approval.

Usage: python preparar_sabana.py AUDIT_DIRECTORY OUTPUT_JSON
Requires the original ZIP sources, inventory.json and dry-run.json from the audit.
"""
import sys, json, zipfile, io, re, datetime, unicodedata
from pathlib import Path
import openpyxl
from openpyxl.utils import get_column_letter

def normal(value):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(value or '')) if not unicodedata.combining(c)).upper().strip()

def number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None

def prepare(root):
    inventory = json.loads((root / 'inventory.json').read_text(encoding='utf-8'))
    dry = json.loads((root / 'dry-run.json').read_text(encoding='utf-8'))
    sources, rates, by_path = [], [], {}
    prefixes = {'2026.zip': '2026', 'IDL.zip': 'IDL', 'complementarios.zip': 'complementarios'}
    for item in inventory:
        key = item['archive'] + '/' + item['path']
        excluded = 'NO USAR' in normal(item['path'])
        source = dict(clave=key, archivo=Path(item['path']).name, carpeta=str(Path(item['path']).parent).replace('\\', '/'),
                      hash_sha256=item['sha256'], hojas=len(item['sheets']), estado='excluida' if excluded else 'pendiente',
                      observaciones='Marcado NO USAR en origen.' if excluded else 'Estructura inventariada; extracción de tarifas pendiente.')
        sources.append(source)
        by_path[prefixes[item['archive']] + '/' + item['path']] = (item, source)

    for item, source in by_path.values():
        if item['archive'] != 'IDL.zip':
            continue
        with zipfile.ZipFile(root / 'sources' / item['archive']) as archive:
            content = archive.read(item['path'])
        values = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(content), data_only=False, read_only=True)
        for sheet in values:
            rows = list(sheet.iter_rows())
            raw_rows = list(formulas[sheet.title].iter_rows())
            # Use detailed tables only; summaries mirror them and would duplicate rates.
            idl = sheet.title in ['BBDD', 'P626', 'IDL MACUL SUB']
            ruth = source['archivo'] == 'TARIFAS SRA RUTH.xlsx'
            if not (idl or ruth):
                continue
            header_row = 2 if idl else 1
            headers = {normal(c.value): c.column for c in rows[header_row - 1] if c.value is not None}
            if not {'CARGO', 'TARIFA'}.issubset(headers):
                continue
            cargo_col = headers['CARGO']
            center_col = headers['CENTRO DE COSTO'] if idl else headers['CD']
            tariff_cols = [c.column for c in rows[header_row - 1] if normal(c.value) == 'TARIFA']
            # In IDL the second TARIFA includes overtime: keep the base and explicit overtime unit only.
            price_cols = [(tariff_cols[0], 'mes', 'Base')]
            extra_cols = [c.column for c in rows[header_row - 1] if normal(c.value).replace(' ', '') in ['TARIFAHHEE50%', 'HHEE50%']]
            if extra_cols:
                price_cols.append((extra_cols[-1], 'hora extra 50%', 'HHEE 50%'))
            for ri, row in enumerate(rows[header_row:], header_row + 1):
                get = lambda col: row[col - 1].value if col <= len(row) else None
                cargo, centro = get(cargo_col), get(center_col)
                if not isinstance(cargo, str) or not isinstance(centro, str) or normal(cargo) == 'CARGO':
                    continue
                if idl:
                    year, period = get(headers['ANO']), get(headers['MES'])
                    month_match = re.fullmatch(r'P0?(\d{1,2})', str(period))
                    date = None
                    if number(year) and month_match and 1 <= int(month_match[1]) <= 12:
                        date = f'{int(year):04d}-{int(month_match[1]):02d}-01'
                    period_text = f'{year} {period}'
                    customer = 'IDL'
                    modes = re.findall(r'\b(EST|SUB)\b', normal(centro))
                    mode = modes[0] if len(set(modes)) == 1 else None
                else:
                    date_value = get(headers['FECHA'])
                    date = date_value.date().isoformat() if isinstance(date_value, datetime.datetime) else None
                    period_text = date or sheet.title
                    # CD is preserved in full; customer is not guessed from its first word.
                    customer = 'Por homologar'
                    mode = str(get(headers['EMPRESA']) or '') or None
                for price_col, unit, variant in price_cols:
                    price = number(get(price_col))
                    if price is None:
                        continue
                    components = []
                    end = tariff_cols[0] if variant == 'Base' else len(row)
                    start = 5 if variant == 'Base' else tariff_cols[0] + 1
                    for col in range(start, end + 1):
                        header = rows[header_row - 1][col - 1].value
                        if header is None:
                            continue
                        val, raw = get(col), raw_rows[ri - 1][col - 1].value
                        components.append(dict(concepto=str(header), valor=val, celda=f'{get_column_letter(col)}{ri}',
                                               formula=raw if isinstance(raw, str) and raw.startswith('=') else None))
                    observation = 'Valor guardado; pendiente de conciliación y aprobación. La fecha de referencia identifica el período de origen, no la vigencia contractual.'
                    if ruth and 'OBSERVACION' in headers and get(headers['OBSERVACION']):
                        observation += ' Nota original: ' + str(get(headers['OBSERVACION']))
                    rates.append(dict(fuente_clave=source['clave'], cliente=customer, centro=centro.strip(), cargo=cargo.strip(),
                                      categoria=None, modalidad=mode, unidad=unit if idl or variant != 'Base' else None, variante=variant, moneda='CLP', precio=price,
                                      periodo=period_text, fecha_referencia=date, hoja=sheet.title, celda=row[price_col - 1].coordinate,
                                      conceptos=components, observaciones=observation))
        source['estado'] = 'parcial'
        source['observaciones'] = 'Lectura de tablas detalladas y tarifas base/HHEE. Resúmenes duplicados y escenarios intermedios no se incorporan. Vigencia pendiente de revisión.'
        values.close()
        formulas.close()

    for file in dry['files']:
        if file['status'] != 'listo':
            continue
        item, source = by_path[file['path']]
        source['estado'] = 'parcial'
        source['observaciones'] = 'Bloques reconocidos por el importador histórico. No acredita cobertura de todas las hojas ni vigencia.'
        for record in file['records']:
            sheet = next(s for s in item['sheets'] if s['name'] == record['sheet'])
            col = re.match(r'[A-Z]+', record['cell'])[0]
            components = []
            for row in sheet['top_rows']:
                price_cell = next((c for c in row if re.match(r'[A-Z]+', c[0])[0] == col), None)
                labels = [c for c in row if isinstance(c[1], str) and not c[1].startswith('=')]
                if price_cell and labels and number(price_cell[2]) is not None:
                    components.append(dict(concepto=labels[0][1], valor=price_cell[2], celda=price_cell[0],
                                           formula=price_cell[1] if isinstance(price_cell[1], str) and price_cell[1].startswith('=') else None))
            rates.append(dict(fuente_clave=source['clave'], cliente=record['customer'], centro=record['center'], cargo=record['cargo'],
                              categoria=None, modalidad=record['mode'], unidad=None, variante='Base', moneda='CLP', precio=record['price'],
                              periodo=record['date'], fecha_referencia=record['date'], hoja=record['sheet'], celda=record['cell'],
                              conceptos=components, observaciones='Antecedente leído por importador histórico; unidad, clasificación y vigencia por confirmar. No se recalcularon fórmulas.'))
    return dict(fuentes=sources, tarifas=rates)

if __name__ == '__main__':
    root, output = Path(sys.argv[1]), Path(sys.argv[2])
    result = prepare(root)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    print(json.dumps({'fuentes': len(result['fuentes']), 'referencias': len(result['tarifas']),
                      'clientes': sorted(set(r['cliente'] for r in result['tarifas']))}, ensure_ascii=False))
