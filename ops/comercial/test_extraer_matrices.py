import unittest
from types import SimpleNamespace
from openpyxl.utils import get_column_letter
from extraer_matrices import extract_sheet, reference_date, source_context


def matrix(cells, columns=4):
    return [[SimpleNamespace(value=cells.get(f'{get_column_letter(c)}{r}'), row=r, column=c,
                             coordinate=f'{get_column_letter(c)}{r}')
             for c in range(1, columns + 1)] for r in range(1, 36)]


class MatrixTests(unittest.TestCase):
    def test_periods_and_invalid_dates(self):
        for text, expected in [('P626', '2026-06-01'), ('P10-2026', '2026-10-01'), ('TARIFA ENERO 2026', '2026-01-01'),
                               ('Cotización 09-10-2026', '2026-10-09'), ('31-02-2026', None), ('P1326', None), ('2026', None)]:
            self.assertEqual(expected, reference_date(text))

    def test_salary_and_hhee_are_not_the_tariff(self):
        cells = {'A4': 'TARIFA ENERO 2026', 'A5': 'OPERARIO', 'A6': 'Sueldo Base', 'B6': 600000,
                 'A11': 'Total Imponible', 'B11': 750000, 'A34': 'PRECIO VENTA', 'B34': 1200000.123456, 'D34': 7000}
        rows = matrix(cells)
        raw = matrix(cells | {'B11': '=B6*1.25', 'B34': '=B11*1.6'})
        result, issues = extract_sheet(rows, raw, 'P626')
        self.assertFalse(issues)
        self.assertEqual(1, len(result))
        self.assertEqual('2026-01-01', result[0]['fecha_referencia'])
        self.assertEqual(1200000.123456, result[0]['precio'])
        self.assertEqual('=B6*1.25', result[0]['conceptos'][1]['formula'])
        self.assertEqual('B34', result[0]['celda'])

    def test_no_invented_date_cargo_or_zero_on_error(self):
        cells = {'A5': 'OPERARIO', 'A6': 'Sueldo Base', 'B6': 500000, 'A34': 'PRECIO VENTA', 'B34': 0}
        rows = matrix(cells)
        self.assertIsNone(extract_sheet(rows, rows, 'P626')[0][0]['fecha_referencia'])
        self.assertEqual(0, extract_sheet(rows, rows, 'P626')[0][0]['precio'])
        for change in [{'A5': 'CARGO'}, {'B34': None}, {'B6': '#REF!'}]:
            rows = matrix(cells | change)
            accepted, issues = extract_sheet(rows, rows, 'P626')
            self.assertEqual([], accepted)
            self.assertEqual(1, len(issues))

    def test_independent_horizontal_blocks_beyond_af(self):
        cells = {'AG4': 'P626', 'AG5': 'SUPERVISOR', 'AG6': 'Sueldo Base', 'AH6': 500000,
                 'AG34': 'PRECIO VENTA', 'AH34': 900000,
                 'AK4': 'P082026', 'AK5': 'APILADOR', 'AK6': 'Sueldo Base', 'AL6': 700000,
                 'AK34': 'PRECIO VENTA', 'AL34': 1100000}
        rows = matrix(cells, columns=38)
        result, _ = extract_sheet(rows, rows, 'Hoja1')
        self.assertEqual(['SUPERVISOR', 'APILADOR'], [x['cargo'] for x in result])
        self.assertEqual(['2026-06-01', '2026-08-01'], [x['fecha_referencia'] for x in result])

    def test_customer_and_mode_do_not_guess_ambiguous_sources(self):
        self.assertEqual(('Maersk', 'MAERSK MALVILLA SUB', 'SUB'), source_context('2026/1-SUB/MAERSK MALVILLA SUB - TARIFAS 2025.xlsx'))
        self.assertEqual(('Por homologar', 'CENCOSUD EST - SUB', None), source_context('2026/2-EST/CENCOSUD EST - SUB - TARIFAS 2025.xlsx'))


if __name__ == '__main__':
    unittest.main()
