"""Read audited client formats, preserving separate blocks and source values."""
import io
import copy
import json
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace
import openpyxl
from preparar_sabana import normal, number
from extraer_matrices import reference_date, source_context
from extraer_tablas import extract_table


def polar(rows, raw, title):
    results = []
    for col in range(len(rows[0])):
        def v(r, offset=0):
            return rows[r-1][col+offset].value if r <= len(rows) and col+offset < len(rows[r-1]) else None
        if normal(v(4)) != 'SUELDO BASE' or normal(v(3, 2)) != 'VALOR' or not normal(v(3)).startswith('ESTRUCTURA DE SUELDO'):
            continue
        cargo = v(2)
        if not isinstance(cargo, str) or not cargo.strip():
            continue
        labels = {' '.join(normal(v(r)).split()): r for r in range(4, len(rows)+1) if isinstance(v(r), str)}
        final = labels.get('TOTAL A FACTURAR FINAL')
        if not final or 'TOTAL NETO CON FEE MENSUAL' not in labels:
            continue
        components = []
        for r in range(4, len(rows)+1):
            if not isinstance(v(r), str):
                continue
            cell = rows[r-1][col+2]
            formula = raw[r-1][col+2].value
            components.append(dict(concepto=v(r).strip(), valor=cell.value, celda=cell.coordinate,
                                   formula=formula if isinstance(formula, str) and formula.startswith('=') else None))
        if any(isinstance(c['valor'], str) and c['valor'].startswith('#') for c in components):
            continue
        for label, unit in [('TOTAL A FACTURAR FINAL', 'mes'), ('HORA ORDINARIA', 'hora'), ('HHEE 50%', 'hora extra 50%'), ('HHEE 100%', 'hora extra 100%')]:
            r = labels.get(label)
            if not r or number(v(r, 2)) is None or v(r, 2) < 0:
                continue
            date = reference_date(v(1))
            results.append(dict(cargo=cargo.strip(), precio=v(r, 2), unidad=unit, variante=label,
                                hoja=title, celda=rows[r-1][col+2].coordinate, conceptos=components,
                                fecha_referencia=date, periodo=str(v(1)) if date else 'Período por confirmar'))
    return results


def parallel_tables(rows, raw, title):
    """Split each header at CARGO and stop at the next header/blank cargo row."""
    results = []
    for ri, row in enumerate(rows):
        starts = [i for i,c in enumerate(row) if normal(c.value) == 'CARGO']
        if not starts:
            continue
        for index, start in enumerate(starts):
            end = starts[index+1] if index+1 < len(starts) else len(row)
            headers = [normal(c.value) for c in row[start:end]]
            if not any(h.startswith('SUELDO BASE') for h in headers):
                continue
            stop = ri+1
            while stop < len(rows) and isinstance(rows[stop][start].value, str) and normal(rows[stop][start].value) not in ('', 'CARGO', 'TOTAL'):
                stop += 1
            block, formulas = [], []
            for local_r, source_r in enumerate(range(ri, stop), 1):
                br, fr = [], []
                for local_c, source_c in enumerate(range(start, end), 1):
                    cell, fcell = rows[source_r][source_c], raw[source_r][source_c]
                    value = cell.value
                    if local_r == 1 and normal(value) == 'PRECIO DE VENTA MES':
                        value = 'TARIFA'
                    if local_r == 1 and normal(value) == 'COSTO FINAL MES' and 'PRECIO DE VENTA MES' in headers:
                        value = 'Costo intermedio (no tarifa)'
                    br.append(SimpleNamespace(value=value, column=local_c, coordinate=f'{openpyxl.utils.get_column_letter(local_c)}{local_r}'))
                    fr.append(SimpleNamespace(value=fcell.value))
                block.append(br)
                formulas.append(fr)
            found, _ = extract_table(block, formulas, title)
            def original(address):
                r,c = openpyxl.utils.cell.coordinate_to_tuple(address)
                return f'{openpyxl.utils.get_column_letter(start+c)}{ri+r}'
            for original_rate in found:
                rate = copy.deepcopy(original_rate)
                rate['celda'] = original(rate['celda'])
                for concept in rate['conceptos']:
                    _, local_c = openpyxl.utils.cell.coordinate_to_tuple(concept['celda'])
                    concept['concepto'] = str(row[start+local_c-1].value).strip()
                    concept['celda'] = original(concept['celda'])
                if rate['variante'] == 'TARIFA' and 'PRECIO DE VENTA MES' in headers:
                    rate['unidad'], rate['variante'] = 'mes', 'PRECIO DE VENTA MES'
                # Do not merge parallel versions with no explicit date.
                rate['variante'] += f' · bloque {row[start].coordinate}'
                rate.pop('centro_fila', None)
                mode = rate.pop('modalidad_fila', None)
                if mode:
                    rate['variante'] += f' · modalidad origen {mode}'
                results.append(rate)
    return results


def extend(root):
    base = json.loads((root/'lote-sabana-tablas.json').read_text(encoding='utf-8'))
    inventory = json.loads((root/'inventory.json').read_text(encoding='utf-8'))
    sources = {s['clave']: s for s in base['fuentes']}
    selected = {'LA POLAR SUB - TARIFAS 2025 FORMATO CLIENTE.xlsx': polar,
                'LA POLAR EST - TARIFAS 2025 FORMATO CLIENTE.xlsx': polar,
                'LTS PEÑON SUB - TARIFAS 2025 REFORMA PREVISIONAL.xlsx': parallel_tables,
                'LTS QUILICURA SUB - TARIFA ADMINISTRATIVO UNIFORME 2025.xlsx': parallel_tables}
    new, coverage = [], []
    for item in inventory:
        reader = selected.get(Path(item['path']).name)
        if not reader:
            continue
        key = item['archive']+'/'+item['path']
        with zipfile.ZipFile(root/'sources'/item['archive']) as z:
            content = z.read(item['path'])
        values = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
        customer, center, mode = source_context(item['path'])
        if reader == parallel_tables:
            customer = 'Walmart Chile'
        count = 0
        for sh in values:
            for rate in reader(list(sh.iter_rows()), list(formulas[sh.title].iter_rows()), sh.title):
                new.append(dict(fuente_clave=key, cliente=customer, centro=center, modalidad=mode,
                                categoria=None, moneda='CLP', **rate,
                                observaciones='Formato cliente auditado. Valores y fórmulas guardados; fecha solo de bloque explícito. Versiones paralelas conservadas, vigencia pendiente de revisión comercial.'))
                count += 1
        values.close()
        formulas.close()
        if count:
            sources[key]['estado'] = 'parcial'
            sources[key]['observaciones'] = 'Formato cliente extraído con conceptos y celdas originales. Versiones, fechas faltantes y vigencia requieren revisión comercial.'
        coverage.append(dict(fuente=key, nuevas=count))
    delta = dict(base, fuentes=[sources[c['fuente']] for c in coverage], tarifas=new)
    base['tarifas'].extend(new)
    for name, content in [('lote-formatos-cliente-delta.json', delta), ('lote-sabana-formatos-cliente.json',base), ('cobertura-formatos-cliente.json',coverage)]:
        (root/name).write_text(json.dumps(content, ensure_ascii=False, default=str), encoding='utf-8')
    print(json.dumps(dict(nuevas=len(new), total=len(base['tarifas']), fuentes=coverage),ensure_ascii=False))


if __name__ == '__main__':
    extend(Path(sys.argv[1]))
